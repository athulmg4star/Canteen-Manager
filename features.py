"""Version 3 workflows. Accounting ledgers remain the source of truth."""
import calendar
import csv
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path
import secrets
import sqlite3
import threading
from flask import abort, g, jsonify, redirect, render_template, request, send_file, session, url_for

# Only these known tables can be audited/undone; identifiers never come from input.
DATED = {
    'Meal_Log_Raw': ('log_id', 'log_date'), 'Payments_Raw': ('payment_id', 'payment_date'),
    'Purchases_Raw': ('purchase_id', 'purchase_date'), 'Misc_Income_Raw': ('income_id', 'income_date'),
    'Staff_Loans_Raw': ('loan_id', 'loan_date'), 'Manager_Fund_Ledger': ('entry_id', 'entry_date'),
    'Consumption_Log': ('log_id', 'purchase_date'), 'Stock_Movements': ('movement_id', 'movement_date'),
    'Service_Days': ('day', 'day'),
}
AUDITED = dict(DATED, Employees=('employee_id', None), Items_Master=('item_id', None),
               Global_Settings=('id', None), Manager_Fund_Settings=('id', None),
               Stock_Items=('item_id', None), Closed_Months=('month', None),
               Expense_Categories=('category_id',None), Vendors=('vendor_id',None),
               Purchase_Descriptions=('description_id',None),Consumption_Items=('item_id',None),
               Price_History=('history_id',None),Prediction_Month_Settings=('month',None))
UNDO_TABLES = {'Meal_Log_Raw', 'Payments_Raw', 'Purchases_Raw', 'Misc_Income_Raw', 'Manager_Fund_Ledger'}


def install_schema(cursor):
    cursor.executescript('''
        CREATE TABLE IF NOT EXISTS Closed_Months (
            month TEXT PRIMARY KEY, closed_at TEXT NOT NULL, note TEXT NOT NULL,
            expected_balance REAL NOT NULL, counted_balance REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS Service_Days (day TEXT PRIMARY KEY, status TEXT NOT NULL CHECK(status IN ('open','closed')), note TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS Stock_Items (item_id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL,
            unit TEXT NOT NULL, reorder_level REAL NOT NULL DEFAULT 0 CHECK(reorder_level>=0));
        CREATE TABLE IF NOT EXISTS Stock_Movements (movement_id INTEGER PRIMARY KEY, item_id INTEGER NOT NULL REFERENCES Stock_Items(item_id),
            movement_date TEXT NOT NULL, kind TEXT NOT NULL CHECK(kind IN ('received','used','wasted')),
            quantity REAL NOT NULL CHECK(quantity>0), note TEXT NOT NULL DEFAULT '');
        CREATE INDEX IF NOT EXISTS stock_item_date ON Stock_Movements(item_id,movement_date);
        CREATE TABLE IF NOT EXISTS Audit_Events (event_id INTEGER PRIMARY KEY, changed_at TEXT NOT NULL,
            table_name TEXT NOT NULL, record_id TEXT NOT NULL, action TEXT NOT NULL, before_json TEXT, after_json TEXT);
        CREATE INDEX IF NOT EXISTS audit_record ON Audit_Events(table_name,record_id,event_id);
        CREATE TABLE IF NOT EXISTS Audit_Undone (event_id INTEGER PRIMARY KEY, undone_at TEXT NOT NULL, reason TEXT NOT NULL);
    ''')
    for table, (pk, date) in AUDITED.items():
        columns = [row[1] for row in cursor.execute(f'PRAGMA table_info("{table}")')]
        if not columns:
            continue
        def row_json(alias):
            return 'json_object(' + ','.join(f"'{c}',{alias}.\"{c}\"" for c in columns) + ')'
        for action, alias in [('INSERT', 'NEW'), ('UPDATE', 'NEW'), ('DELETE', 'OLD')]:
            before = row_json('OLD') if action != 'INSERT' else 'NULL'
            after = row_json('NEW') if action != 'DELETE' else 'NULL'
            change = f'WHEN {before} IS NOT {after}' if action == 'UPDATE' else ''
            cursor.execute(f'''CREATE TRIGGER IF NOT EXISTS audit_v3_{table}_{action} AFTER {action} ON "{table}" {change}
                BEGIN INSERT INTO Audit_Events(changed_at,table_name,record_id,action,before_json,after_json)
                VALUES(strftime('%Y-%m-%d %H:%M:%f','now','localtime'),'{table}',{alias}."{pk}",'{action}',{before},{after}); END''')
            if date:
                aliases = ['NEW'] if action == 'INSERT' else ['OLD'] if action == 'DELETE' else ['OLD','NEW']
                # A backdated write can change the opening balance of later closed months too.
                conditions = ' OR '.join(f'EXISTS(SELECT 1 FROM Closed_Months WHERE month >= substr({a}."{date}",1,7))' for a in aliases)
                cursor.execute(f'''CREATE TRIGGER IF NOT EXISTS locked_v3_{table}_{action} BEFORE {action} ON "{table}"
                    WHEN {conditions} BEGIN SELECT RAISE(ABORT,'This date affects a closed month. Reopen the affected months first.'); END''')
    for table, column in [('Employees','opening_balance'),('Global_Settings','canteen_opening_balance'),('Manager_Fund_Settings','opening_balance')]:
        cursor.execute(f'''CREATE TRIGGER IF NOT EXISTS lock_opening_v3_{table} BEFORE UPDATE OF {column} ON {table}
          WHEN OLD.{column} IS NOT NEW.{column} AND EXISTS(SELECT 1 FROM Closed_Months)
          BEGIN SELECT RAISE(ABORT,'Opening balances affect closed months. Reopen them before editing.'); END''')
    cursor.executescript('''
        CREATE TRIGGER IF NOT EXISTS stock_available_v3 BEFORE INSERT ON Stock_Movements
        WHEN NEW.kind != 'received' AND NEW.quantity > COALESCE((SELECT SUM(CASE WHEN kind='received' THEN quantity ELSE -quantity END) FROM Stock_Movements WHERE item_id=NEW.item_id),0)+0.0000001
        BEGIN SELECT RAISE(ABORT,'Not enough stock for this movement.'); END;
    ''')

    cursor.executescript("""
        CREATE TRIGGER IF NOT EXISTS stock_chronology_v3 BEFORE INSERT ON Stock_Movements
        WHEN NEW.kind != 'received' AND (
          NEW.quantity > COALESCE((SELECT SUM(CASE WHEN kind='received' THEN quantity ELSE -quantity END) FROM Stock_Movements WHERE item_id=NEW.item_id AND movement_date<=NEW.movement_date),0)+0.0000001
          OR NEW.quantity > COALESCE((SELECT MIN(balance) FROM (
            SELECT movement_date,SUM(CASE WHEN kind='received' THEN quantity ELSE -quantity END)
            OVER (ORDER BY movement_date,movement_id) AS balance FROM Stock_Movements WHERE item_id=NEW.item_id
          ) WHERE movement_date>NEW.movement_date),1e100)+0.0000001)
        BEGIN SELECT RAISE(ABORT,'Not enough stock on this date or a later recorded date.'); END;
        CREATE TRIGGER IF NOT EXISTS employee_opening_delete_v3 BEFORE DELETE ON Employees
        WHEN COALESCE(OLD.opening_balance,0)!=0 AND EXISTS(SELECT 1 FROM Closed_Months)
        BEGIN SELECT RAISE(ABORT,'Opening balances affect closed months. Reopen them before editing.'); END;
        CREATE TRIGGER IF NOT EXISTS employee_opening_insert_v3 BEFORE INSERT ON Employees
        WHEN COALESCE(NEW.opening_balance,0)!=0 AND EXISTS(SELECT 1 FROM Closed_Months)
        BEGIN SELECT RAISE(ABORT,'Opening balances affect closed months. Reopen them before editing.'); END;
    """)


def month_bounds(value=None):
    value = value or dt.date.today().strftime('%Y-%m')
    try:
        start = dt.datetime.strptime(value, '%Y-%m').date()
        if start.strftime('%Y-%m') != value:
            raise ValueError()
    except (ValueError, TypeError):
        abort(400, description='Choose a valid month.')
    last = start.replace(day=calendar.monthrange(start.year, start.month)[1])
    return value, start, last


def finite(value, minimum=None):
    try:
        n = float(value)
        if not math.isfinite(n) or (minimum is not None and n < minimum):
            raise ValueError()
        return n
    except (ValueError, TypeError):
        abort(400, description='Enter a valid number within the allowed range.')


def records(conn, sql, args=()):
    cursor = conn.execute(sql, args)
    return [dict(zip([c[0] for c in cursor.description], row)) for row in cursor]


def scalar(conn, sql, args=()):
    return conn.execute(sql, args).fetchone()[0] or 0


def balances(conn, month):
    _, first, last = month_bounds(month)
    opening = scalar(conn, 'SELECT canteen_opening_balance FROM Global_Settings WHERE id=1')
    def total(table, date, before=None):
        return scalar(conn, f'SELECT COALESCE(SUM(amount),0) FROM {table} WHERE {date} < ?', (str(before),))
    initial = opening + total('Payments','payment_date',first) + total('Misc_Income','income_date',first) - total('Purchases','purchase_date',first)
    end = last + dt.timedelta(days=1)
    closing = opening + total('Payments','payment_date',end) + total('Misc_Income','income_date',end) - total('Purchases','purchase_date',end)
    income = scalar(conn, 'SELECT SUM(amount) FROM Payments WHERE payment_date BETWEEN ? AND ?', (str(first),str(last)))
    misc = scalar(conn, 'SELECT SUM(amount) FROM Misc_Income WHERE income_date BETWEEN ? AND ?', (str(first),str(last)))
    expense = scalar(conn, 'SELECT SUM(amount) FROM Purchases WHERE purchase_date BETWEEN ? AND ?', (str(first),str(last)))
    meals = scalar(conn, 'SELECT SUM(quantity*locked_price) FROM Meal_Log WHERE log_date BETWEEN ? AND ?', (str(first),str(last)))
    return dict(opening=initial, closing=closing, payments=income, misc=misc, income=income+misc, expense=expense, meals=meals)


def stock_rows(conn):
    return records(conn, '''SELECT s.item_id,s.name,s.unit,s.reorder_level,
        COALESCE(SUM(CASE WHEN m.kind='received' THEN m.quantity ELSE -m.quantity END),0) AS balance
        FROM Stock_Items s LEFT JOIN Stock_Movements m ON m.item_id=s.item_id GROUP BY s.item_id ORDER BY s.name''')


def day_rows(conn, month):
    _, first, last = month_bounds(month)
    logs = {r['day']:r for r in records(conn, '''SELECT log_date AS day, COUNT(DISTINCT employee_id) AS staff,
            SUM(quantity) AS portions, SUM(quantity*locked_price) AS amount FROM Meal_Log
            WHERE log_date BETWEEN ? AND ? GROUP BY log_date''', (str(first),str(last)))}
    status = {r['day']:r for r in records(conn,'SELECT * FROM Service_Days WHERE day BETWEEN ? AND ?', (str(first),str(last)))}
    result=[]
    for number in range(1,last.day+1):
        day=first.replace(day=number); key=str(day)
        expected=status.get(key,{}).get('status', 'open' if day.weekday()<5 else 'closed')=='open'
        result.append(dict(day=key,number=number,weekday=day.strftime('%a'),expected=expected,
            missing=expected and day<=dt.date.today() and key not in logs, future=day>dt.date.today(),
            note=status.get(key,{}).get('note',''),**{k:v for k,v in logs.get(key,dict(staff=0,portions=0,amount=0)).items() if k!='day'}))
    return result


REPORTS = {'meals':'Meal ledger','payments':'Staff payments','expenses':'Purchases','misc':'Other income','stock':'Stock movements'}

def report_rows(conn, kind, month, staff='', search=''):
    _, first,last=month_bounds(month)
    queries={
      'meals':'''SELECT m.log_date AS Date,e.name AS Staff,m.item_name AS Detail,m.quantity AS Quantity,m.locked_price AS Rate,m.quantity*m.locked_price AS Amount FROM Meal_Log m JOIN Employees e ON e.employee_id=m.employee_id WHERE m.log_date BETWEEN ? AND ?''',
      'payments':'''SELECT p.payment_date AS Date,e.name AS Staff,'Payment' AS Detail,p.amount AS Amount FROM Payments p JOIN Employees e ON e.employee_id=p.employee_id WHERE p.payment_date BETWEEN ? AND ?''',
      'expenses':'''SELECT purchase_date AS Date,vendor AS Vendor,item_category AS Category,description AS Detail,amount AS Amount FROM Purchases WHERE purchase_date BETWEEN ? AND ?''',
      'misc':'''SELECT income_date AS Date,source AS Detail,amount AS Amount FROM Misc_Income WHERE income_date BETWEEN ? AND ?''',
      'stock':'''SELECT m.movement_date AS Date,s.name AS Item,m.kind AS Movement,m.quantity AS Quantity,s.unit AS Unit,m.note AS Detail FROM Stock_Movements m JOIN Stock_Items s ON s.item_id=m.item_id WHERE m.movement_date BETWEEN ? AND ?''',
    }
    if kind not in queries:
        abort(400,description='Choose an available report.')
    query=queries[kind]; args=[str(first),str(last)]
    if staff and kind in ('meals','payments'):
        query+=' AND e.employee_id=?';args.append(staff)
    rows=records(conn,query+' ORDER BY Date',args)
    if search:
        rows=[r for r in rows if search.casefold() in ' '.join(str(v) for v in r.values()).casefold()]
    return rows


def register(app, connect, db_path):
    # Serialize local requests during database restore and multi-step workflows.
    mutex=threading.RLock()
    def acquire():
        if request.endpoint!='static':
            mutex.acquire();g.feature_lock=True
    app.before_request_funcs.setdefault(None,[]).insert(0,acquire)
    @app.teardown_request
    def release(error=None):
        for connection in g.pop('db_connections',[]):
            connection.close()
        if g.pop('feature_lock',False):
            mutex.release()
    @app.errorhandler(sqlite3.IntegrityError)
    def database_guard(error):
        message=str(error)
        if not any(t in message for t in ('closed month','closed months','Not enough stock')):
            message='This change conflicts with an existing record. Check the entry and try again.'
        return render_template('error.html',message=message),409

    @app.context_processor
    def feature_context():
        return dict(today=dt.date.today().isoformat(),current_month=dt.date.today().strftime('%Y-%m'))

    def page(template, **context):
        return render_template('features/'+template+'.html',**context)

    def overview_payload(month):
        with connect() as conn:
            totals=balances(conn,month)
            _,first,_=month_bounds(month)
            previous=(first-dt.timedelta(days=1)).strftime('%Y-%m')
            days=day_rows(conn,month)
            stock=stock_rows(conn)
            return dict(month=month,totals=totals,previous=balances(conn,previous),previous_month=previous,
                days=days,low_stock=[r for r in stock if r['balance']<=r['reorder_level']],
                closed=bool(conn.execute('SELECT 1 FROM Closed_Months WHERE month=?',(month,)).fetchone()))

    @app.get('/overview')
    def overview():
        month,_,_=month_bounds(request.args.get('month'))
        return page('overview',title='Canteen overview',data=overview_payload(month))
    # Keep the original charts accessible while making the new overview the home screen.
    app.add_url_rule('/classic-dashboard','classic_dashboard',app.view_functions['dashboard'])
    app.view_functions['dashboard']=overview

    @app.get('/api/overview')
    def overview_api():
        month,_,_=month_bounds(request.args.get('month'))
        return jsonify(overview_payload(month))

    @app.get('/api/commands')
    def commands():
        with connect() as conn:
            staff=records(conn,'SELECT employee_id,name,department FROM Employees ORDER BY name')
        return jsonify([dict(label=r['name'],hint=r['department'] or 'Staff',url=f"/timeline/{r['employee_id']}") for r in staff])

    @app.route('/calendar',methods=['GET','POST'])
    def service_calendar():
        if request.method=='POST':
            day=request.form.get('day',''); status=request.form.get('status')
            try: dt.date.fromisoformat(day)
            except ValueError: abort(400,description='Choose a valid calendar date.')
            if status not in ('open','closed'):abort(400,description='Choose open or closed.')
            with connect() as conn:
                conn.execute('INSERT INTO Service_Days VALUES (?,?,?) ON CONFLICT(day) DO UPDATE SET status=excluded.status,note=excluded.note', (day,status,request.form.get('note','').strip()))
            return redirect(url_for('service_calendar',month=day[:7],success=1))
        month,first,last=month_bounds(request.args.get('month'))
        with connect() as conn: days=day_rows(conn,month)
        return page('calendar',title='Meal calendar',month=month,days=days,padding=first.weekday())

    @app.route('/inventory',methods=['GET','POST'])
    def inventory():
        if request.method=='POST':
            with connect() as conn:
                if request.form.get('action')=='item':
                    name=request.form.get('name','').strip();unit=request.form.get('unit','').strip()
                    if not name or not unit:abort(400,description='Give the item a name and unit.')
                    threshold=finite(request.form.get('reorder_level'),0)
                    conn.execute('INSERT INTO Stock_Items(name,unit,reorder_level) VALUES (?,?,?)',(name,unit,threshold))
                elif request.form.get('action')=='threshold':
                    conn.execute('UPDATE Stock_Items SET reorder_level=? WHERE item_id=?',(finite(request.form.get('reorder_level'),0),request.form.get('item_id')))
                else:
                    item=request.form.get('item_id');kind=request.form.get('kind');quantity=finite(request.form.get('quantity'),0.000001)
                    date=request.form.get('movement_date','')
                    try:dt.date.fromisoformat(date)
                    except ValueError:abort(400,description='Choose a movement date.')
                    if kind not in ('received','used','wasted') or not conn.execute('SELECT 1 FROM Stock_Items WHERE item_id=?',(item,)).fetchone():
                        abort(400,description='Choose an existing item and movement type.')
                    conn.execute('INSERT INTO Stock_Movements(item_id,movement_date,kind,quantity,note) VALUES (?,?,?,?,?)',(item,date,kind,quantity,request.form.get('note','').strip()))
            return redirect('/inventory?success=1')
        with connect() as conn:
            items=stock_rows(conn)
            movements=records(conn,'''SELECT m.*,s.name,s.unit FROM Stock_Movements m JOIN Stock_Items s USING(item_id) ORDER BY m.movement_date DESC,m.movement_id DESC LIMIT 100''')
        return page('inventory',title='Inventory',items=items,movements=movements)

    @app.route('/closing',methods=['GET','POST'])
    def closing():
        month,first,last=month_bounds(request.form.get('month') if request.method=='POST' else request.args.get('month'))
        with connect() as conn:
            totals=balances(conn,month);days=day_rows(conn,month)
            if request.method=='POST':
                note=request.form.get('note','').strip()
                if not note:abort(400,description='Add a note explaining this close or reopen action.')
                if request.form.get('action')=='reopen':
                    conn.execute('DELETE FROM Closed_Months WHERE month=?',(month,))
                    conn.execute("INSERT INTO Audit_Events(changed_at,table_name,record_id,action,after_json) VALUES (datetime('now','localtime'),'Closed_Months',?,'REOPEN',?)",(month,json.dumps({'reason':note})))
                else:
                    if last>dt.date.today():abort(409,description='Close this month after its final day.')
                    if any(request.form.get(k)!='yes' for k in ('review_days','review_dues','review_reports')):
                        abort(400,description='Complete all three review checks.')
                    counted=finite(request.form.get('counted_balance'))
                    if abs(counted-totals['closing'])>0.005:
                        abort(409,description='The counted balance does not match the ledger. Reconcile the difference before closing.')
                    conn.execute('INSERT INTO Closed_Months VALUES (?,?,?,?,?)',(month,dt.datetime.now().isoformat(timespec='seconds'),note,totals['closing'],counted))
                return redirect(url_for('closing',month=month,success=1))
            closed=records(conn,'SELECT * FROM Closed_Months WHERE month=?',(month,))
            dues=records(conn,'''SELECT e.name, e.opening_balance + COALESCE((SELECT SUM(quantity*locked_price) FROM Meal_Log WHERE employee_id=e.employee_id AND log_date<=?),0)
                - COALESCE((SELECT SUM(amount) FROM Payments WHERE employee_id=e.employee_id AND payment_date<=?),0) AS due FROM Employees e ORDER BY due DESC''',(str(last),str(last)))
        return page('closing',title='Month closing',month=month,totals=totals,missing=[d for d in days if d['missing']],dues=[r for r in dues if r['due']>0.005],closed=closed[0] if closed else None)

    @app.get('/timeline/<int:emp_id>')
    def timeline(emp_id):
        with connect() as conn:
            staff=records(conn,'SELECT * FROM Employees WHERE employee_id=?',(emp_id,))
            if not staff:abort(404)
            employee=staff[0]
            entries=records(conn,'''SELECT log_date AS day,'Meal' AS kind,item_name AS detail,quantity*locked_price AS debit,0 AS credit,log_id AS id FROM Meal_Log WHERE employee_id=?
                UNION ALL SELECT payment_date,'Payment','Staff payment',0,amount,payment_id FROM Payments WHERE employee_id=?
                ORDER BY day,id''',(emp_id,emp_id))
            balance=employee['opening_balance'] or 0
            groups={}
            for row in entries:
                balance+=row['debit']-row['credit'];row['balance']=balance
                groups.setdefault(row['day'][:7],[]).append(row)
            loans=records(conn,'SELECT * FROM Staff_Loans WHERE employee_id=? ORDER BY loan_date DESC',(emp_id,))
            changes=records(conn,'''SELECT * FROM Audit_Events WHERE table_name IN ('Meal_Log_Raw','Payments_Raw','Staff_Loans_Raw','Employees')
              AND COALESCE(json_extract(after_json,'$.employee_id'),json_extract(before_json,'$.employee_id'))=? ORDER BY event_id DESC LIMIT 50''',(emp_id,))
        return page('timeline',title=employee['name']+' · account',employee=employee,groups=sorted(groups.items(),reverse=True),balance=balance,loans=loans,changes=changes)

    @app.get('/audit')
    def audit():
        before=request.args.get('before',type=int) or 9223372036854775807
        with connect() as conn:
            events=records(conn,'''SELECT a.*,u.undone_at,u.reason AS undo_reason FROM Audit_Events a LEFT JOIN Audit_Undone u USING(event_id) WHERE a.event_id<? ORDER BY a.event_id DESC LIMIT 100''',(before,))
            for event in events:
                event['before']=json.loads(event['before_json']) if event['before_json'] else None
                event['after']=json.loads(event['after_json']) if event['after_json'] else None
                event['can_undo']=event['table_name'] in UNDO_TABLES and not event['undone_at'] and event['action'] in ('INSERT','UPDATE')
        return page('audit',title='Change history',events=events)

    @app.post('/audit/undo')
    def undo():
        reason=request.form.get('reason','').strip()
        if not reason:abort(400,description='Add a reason for undoing the change.')
        with connect() as conn:
            conn.execute('BEGIN IMMEDIATE')
            found=records(conn,'SELECT * FROM Audit_Events WHERE event_id=?',(request.form.get('event_id'),))
            if not found:abort(404)
            event=found[0];table=event['table_name']
            if table not in UNDO_TABLES or event['action'] not in ('INSERT','UPDATE'):abort(409,description='Use the original workflow or Recycle Bin to reverse this change.')
            if conn.execute('SELECT 1 FROM Audit_Undone WHERE event_id=?',(event['event_id'],)).fetchone() or conn.execute('SELECT 1 FROM Audit_Events WHERE table_name=? AND record_id=? AND event_id>?',(table,event['record_id'],event['event_id'])).fetchone():
                abort(409,description='This record has changed since then. Review its latest change first.')
            pk=AUDITED[table][0]
            if table=='Payments_Raw' and conn.execute('SELECT 1 FROM Staff_Loans_Raw WHERE linked_payment_id=?',(event['record_id'],)).fetchone():
                abort(409,description='Manage linked payments through Staff Loans.')
            current=records(conn,f'SELECT * FROM "{table}" WHERE "{pk}"=?',(event['record_id'],))
            if not current or current[0]!=json.loads(event['after_json']):abort(409,description='The record no longer matches this change.')
            if event['action']=='INSERT':
                if 'is_deleted' in current[0]:conn.execute(f'UPDATE "{table}" SET is_deleted=1,deleted_at=? WHERE "{pk}"=?',(dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),event['record_id']))
                else:conn.execute(f'DELETE FROM "{table}" WHERE "{pk}"=?',(event['record_id'],))
            else:
                old=json.loads(event['before_json']);columns=[c for c in old if c!=pk]
                conn.execute(f'UPDATE "{table}" SET '+','.join(f'"{c}"=?' for c in columns)+f' WHERE "{pk}"=?',[old[c] for c in columns]+[event['record_id']])
            conn.execute('INSERT INTO Audit_Undone VALUES (?,?,?)',(event['event_id'],dt.datetime.now().isoformat(timespec='seconds'),reason))
        return redirect('/audit?success=1')

    @app.get('/reports')
    def reports():
        month,_,_=month_bounds(request.args.get('month'));kind=request.args.get('kind','meals')
        staff=request.args.get('staff','');search=request.args.get('search','').strip()
        with connect() as conn:
            rows=report_rows(conn,kind,month,staff,search)
            employees=records(conn,'SELECT employee_id,name FROM Employees ORDER BY name')
        columns=list(rows[0]) if rows else []
        export=request.args.get('format')
        filename=f'{kind}_{month}'
        if export=='csv':
            buffer=io.StringIO();writer=csv.writer(buffer);writer.writerow(columns)
            for row in rows:
                writer.writerow([("'"+str(row[c])) if isinstance(row[c],str) and row[c].startswith(('=','+','-','@')) else row[c] for c in columns])
            return send_file(io.BytesIO(buffer.getvalue().encode('utf-8-sig')),mimetype='text/csv',as_attachment=True,download_name=filename+'.csv')
        if export=='xlsx':
            from openpyxl import Workbook
            book=Workbook();sheet=book.active;sheet.title='Report';sheet.append(columns)
            for row in rows:sheet.append(list(row.values()))
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value,str):cell.data_type='s'
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in column)+2))
            out=io.BytesIO();book.save(out);out.seek(0)
            return send_file(out,as_attachment=True,download_name=filename+'.xlsx')
        if export=='pdf':
            from fpdf import FPDF
            pdf=FPDF(orientation='L');pdf.set_auto_page_break(True,15);pdf.add_page();pdf.set_font('Helvetica',size=14)
            pdf.cell(0,10,text=f'{REPORTS[kind]} | {month}',new_x='LMARGIN',new_y='NEXT')
            pdf.set_font('Helvetica',size=8)
            # Table repeats its heading and wraps long text across rows/pages.
            if rows:
                with pdf.table(col_widths=tuple(1 for _ in columns),line_height=5) as table:
                    for values in [columns]+[list(row.values()) for row in rows]:
                        line=table.row()
                        for value in values:line.cell(str(value if value is not None else '').encode('latin-1','replace').decode('latin-1'))
            else:pdf.cell(0,10,text='No records match these filters.')
            return send_file(io.BytesIO(bytes(pdf.output())),mimetype='application/pdf',as_attachment=True,download_name=filename+'.pdf')
        if export:abort(400,description='Choose CSV, Excel or PDF.')
        return page('reports',title='Report centre',month=month,kind=kind,kinds=REPORTS,staff=staff,search=search,employees=employees,rows=rows,columns=columns,total=sum(r.get('Amount',0) or 0 for r in rows))

    @app.get('/api/drilldown')
    def drilldown():
        month,_,_=month_bounds(request.args.get('month'));kind=request.args.get('kind','payments')
        with connect() as conn: rows=report_rows(conn,kind,month)
        return jsonify(rows=rows,total=sum(r.get('Amount',0) or 0 for r in rows))

    @app.get('/preferences')
    def preferences():return page('preferences',title='Appearance & shortcuts')

    def backup_directory():
        directory=Path(db_path()).resolve().parent/'backups';directory.mkdir(exist_ok=True)
        return directory
    def fingerprint(path):return hashlib.sha256(path.read_bytes()).hexdigest()
    def inspect_backup(path):
        with sqlite3.connect(path.as_uri()+'?mode=ro',uri=True) as source:
            if source.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise ValueError('Integrity check failed')
            required={'Employees','Meal_Log_Raw','Payments_Raw','Purchases_Raw','Audit_Events','Closed_Months','Stock_Items','Form_Submissions'}
            actual={r[0] for r in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if not required.issubset(actual):raise ValueError('This snapshot predates the feature upgrade; use the matching application version to restore it.')
            with connect() as current:
                for table in required:
                    expected=[(r[1],r[2]) for r in current.execute(f'PRAGMA table_info("{table}")')]
                    observed=[(r[1],r[2]) for r in source.execute(f'PRAGMA table_info("{table}")')]
                    if observed!=expected:raise ValueError('Snapshot schema does not match this application version.')
                expected_triggers=dict(current.execute("SELECT name,sql FROM sqlite_master WHERE type='trigger' AND (name LIKE '%v3%')"))
                observed_triggers=dict(source.execute("SELECT name,sql FROM sqlite_master WHERE type='trigger' AND (name LIKE '%v3%')"))
                if expected_triggers!=observed_triggers:raise ValueError('Snapshot protections differ from this application version.')
            counts={t:scalar(source,f'SELECT COUNT(*) FROM "{t}"') for t in sorted(required)}
        return dict(name=path.name,bytes=path.stat().st_size,modified=dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),sha256=fingerprint(path),counts=counts)
    def snapshot(prefix='verified'):
        path=backup_directory()/f'{prefix}_{dt.datetime.now():%Y-%m-%d_%H-%M-%S}_{secrets.token_hex(4)}.db'
        with connect() as source,sqlite3.connect(path) as dest:source.backup(dest)
        inspect_backup(path)
        return path
    def selected_backup(name):
        paths={p.name:p for p in backup_directory().glob('*.db') if not p.is_symlink()}
        if name not in paths:abort(404)
        return paths[name]

    @app.route('/backups',methods=['GET','POST'])
    def backups():
        if request.method=='POST':
            path=snapshot()
            return redirect(url_for('backups',selected=path.name,success=1))
        entries=[]
        for path in sorted(backup_directory().glob('*.db'),key=lambda p:p.stat().st_mtime,reverse=True):
            if path.is_symlink():continue
            try:entries.append(dict(inspect_backup(path),valid=True))
            except (ValueError,sqlite3.DatabaseError) as error:entries.append(dict(name=path.name,valid=False,error=str(error)))
        selected=None
        if request.args.get('selected'):
            try:selected=inspect_backup(selected_backup(request.args['selected']))
            except (ValueError,sqlite3.DatabaseError):abort(400,description='This backup is not compatible or failed verification.')
        newest=next((r for r in entries if r['valid']),None)
        overdue=not newest or (dt.datetime.now()-dt.datetime.fromisoformat(newest['modified'])).total_seconds()>7*86400
        with connect() as conn: current={t:scalar(conn,f'SELECT COUNT(*) FROM "{t}"') for t in selected['counts']} if selected else {}
        return page('backups',title='Backup health & recovery',entries=entries,selected=selected,current=current,overdue=overdue)

    @app.get('/backups/download')
    def backup_download():
        path=selected_backup(request.args.get('name'))
        return send_file(path,as_attachment=True,download_name=path.name)

    @app.post('/backups/restore')
    def restore_backup():
        if request.form.get('confirmation')!='RESTORE':abort(400,description='Type RESTORE to confirm replacement with the selected snapshot.')
        path=selected_backup(request.form.get('name'))
        try:info=inspect_backup(path)
        except (ValueError,sqlite3.DatabaseError):abort(400,description='This backup failed verification.')
        if info['sha256']!=request.form.get('sha256'):abort(409,description='This backup changed since the preview. Review it again.')
        safety=snapshot('before_restore')
        # Close all request handles before replacing the database contents.
        for connection in g.pop('db_connections',[]):connection.close()
        try:
            with sqlite3.connect(path) as source,sqlite3.connect(db_path()) as destination:
                source.backup(destination)
                destination.execute('INSERT OR IGNORE INTO Form_Submissions VALUES (?,?)',(g.submission_token,request.path))
                destination.execute("INSERT INTO Audit_Events(changed_at,table_name,record_id,action,after_json) VALUES (datetime('now','localtime'),'Database','restore','RESTORE',?)",(json.dumps(dict(snapshot=path.name,safety_backup=safety.name)),))
        except Exception:
            with sqlite3.connect(safety) as source,sqlite3.connect(db_path()) as destination:source.backup(destination)
            raise
        return redirect('/backups?success=1')
